"""P5: AcroForm fill + flatten (read back with two independent readers), Word write-back keeps formatting, flat overlay."""

from __future__ import annotations

import docx
import pymupdf as fitz
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.forms.acroform import parse_acroform
from app.forms.docx_form import parse_docx_form
from app.forms.flat_pdf import parse_flat_pdf
from app.writeback.acroform import fill_acroform
from app.writeback.docx_form import fill_docx
from app.writeback.flat_pdf import fill_flat_pdf

from .conftest import FIXTURES


def test_acroform_fill_and_flatten(tmp_path):
    src = FIXTURES / "01_acroform_basic" / "form.pdf"
    fields = parse_acroform(src)
    by = {f.key: f for f in fields}
    values = {
        "full_name": "Priya Natarajan",
        "date_of_birth_dd_mm_yyyy": "1990-03-14",
        "marital_status": "De facto",
        "i_would_like_to_receive_the_client_newsletter": True,
        "total_income": "93784.37",
        "signature_of_applicant": "SHOULD NOT BE WRITTEN",
    }
    out = tmp_path / "filled.pdf"
    warnings = fill_acroform(src, out, fields, values)
    assert warnings == []
    # reader 1: PyMuPDF
    doc = fitz.open(out)
    got = {w.field_name: w.field_value for p in doc for w in p.widgets()}
    assert got[by["full_name"].source_ref] == "Priya Natarajan"
    assert got[by["date_of_birth_dd_mm_yyyy"].source_ref] == "14/03/1990"
    assert got[by["total_income"].source_ref] == "93,784.37"
    assert got[by["signature_of_applicant"].source_ref] in ("", None)
    # reader 2: pypdf
    reader = PdfReader(str(out))
    pf = reader.get_fields()
    assert pf[by["full_name"].source_ref]["/V"] == "Priya Natarajan"
    assert pf[by["marital_status"].source_ref]["/V"] == "/Defacto"
    assert pf[by["i_would_like_to_receive_the_client_newsletter"].source_ref]["/V"] not in ("/Off", None)
    # appearance streams exist for the filled text widget
    w = next(w for p in doc for w in p.widgets() if w.field_name == by["full_name"].source_ref)
    assert doc.xref_get_key(w.xref, "AP")[0] == "dict"

    flat = tmp_path / "flat.pdf"
    fill_acroform(src, flat, fields, values, flatten=True)
    fd = fitz.open(flat)
    assert not any(True for p in fd for _ in p.widgets())
    assert "Priya Natarajan" in fd[0].get_text()


def test_docx_writeback_preserves_formatting(tmp_path):
    src = FIXTURES / "02_docx_form" / "form.docx"
    fields = parse_docx_form(src)
    by = {f.label: f for f in fields}
    values = {
        by["Full name"].key: "Amara Okafor",
        by["Date of birth"].key: "1996-11-07",
        by["Preferred contact method"].key: "Email",
        by["Email"].key: "amara.okafor@example.org",
        by["Days available per week"].key: "3",
        by["Emergency contact name"].key: "Chidi Okafor",
        by["Relationship"].key: "Brother",
        by["Do you hold a current driver licence?"].key: "Yes",
        by["Police check consent"].key: True,
    }
    out = tmp_path / "filled.docx"
    assert fill_docx(src, out, fields, values) == []
    d = docx.Document(str(out))
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "Days available per week: 3" in body_text
    assert "Emergency contact name: Chidi Okafor   Relationship: Brother" in body_text
    assert "☒ Yes   ☐ No" in body_text
    xml = d.element.xml
    assert "Amara Okafor" in xml and "07/11/1996" in xml and "showingPlcHdr" not in xml.split("Police check consent")[0][-2000:]
    assert d.tables[0].rows[0].cells[1].text == "amara.okafor@example.org"
    # the bold label run is untouched; the value run in the sdt keeps its (non-placeholder) formatting
    label_run = next(p for p in d.paragraphs if p.text.startswith("Full name")).runs[0]
    assert label_run.bold and label_run.text == "Full name: "
    sdts = list(d.element.body.iter(qn("w:sdt")))
    from lxml import etree
    assert b"PlaceholderText" not in etree.tostring(sdts[0])
    checkbox = sdts[3]
    assert checkbox.find(".//{http://schemas.microsoft.com/office/word/2010/wordml}checked").get(
        "{http://schemas.microsoft.com/office/word/2010/wordml}val") == "1"


def test_flat_overlay(tmp_path):
    src = FIXTURES / "07_flat_pdf" / "form.pdf"
    fields = parse_flat_pdf(src)
    by = {f.label: f for f in fields}
    long_branch = "Glenroy " * 20
    values = {by["Full name"].key: "Lucas Ferreira", by["Send me the monthly events newsletter"].key: True,
              by["Preferred branch"].key: long_branch}
    out = tmp_path / "filled.pdf"
    warnings = fill_flat_pdf(src, out, fields, values)
    text = fitz.open(out)[0].get_text()
    assert "Lucas Ferreira" in text and "X" in text
    assert any("overflows" in w for w in warnings)
