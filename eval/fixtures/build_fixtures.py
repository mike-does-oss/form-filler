"""Build the synthetic evaluation fixtures (no real personal data).

Each set directory gets:
  form.pdf | form.docx       the blank form
  sources/                   source documents
  gold.json                  expected values (+ judgment facts, missing facts, user answers)
  expected_schema.json       hand-written expected field schema (labels, types, options, sections)

Run:  uv run python eval/fixtures/build_fixtures.py
"""

from __future__ import annotations

import io
import json
import random
import shutil
from pathlib import Path

import docx
import pymupdf as fitz
from docx.oxml import parse_xml
from docx.shared import Pt
from PIL import Image, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).parent
STYLES = getSampleStyleSheet()
W, H = A4


# ============================================================ helpers
class AcroFormBuilder:
    """Draws a simple fillable form with reportlab. Widget names are deliberately meaningless."""

    def __init__(self, path: Path, title: str):
        self.c = canvas.Canvas(str(path), pagesize=A4)
        self.form = self.c.acroForm
        self.y = H - 60
        self.n_text = 0
        self.n_check = 0
        self.n_group = 0
        self.c.setFont("Helvetica-Bold", 16)
        self.c.drawString(50, self.y, title)
        self.y -= 30

    def _need(self, h: float) -> None:
        if self.y - h < 60:
            self.new_page()

    def new_page(self) -> None:
        self.c.showPage()
        self.y = H - 60

    def heading(self, text: str) -> None:
        self._need(40)
        self.y -= 8
        self.c.setFont("Helvetica-Bold", 12.5)
        self.c.drawString(50, self.y, text)
        self.y -= 22

    def text(self, label: str, width: float = 300, multiline: bool = False, note: str | None = None) -> str:
        h = 40 if multiline else 18
        self._need(h + 16)
        self.n_text += 1
        name = f"Text{self.n_text}"
        self.c.setFont("Helvetica", 10)
        self.c.drawString(50, self.y - 12, label)
        self.form.textfield(
            name=name, x=230, y=self.y - h + 2, width=width, height=h, borderStyle="inset",
            borderColor=colors.grey, fillColor=colors.white, forceBorder=True, fontSize=10,
            fieldFlags="multiline" if multiline else "",
        )
        self.y -= h + 10
        if note:
            self.c.setFont("Helvetica-Oblique", 8)
            self.c.drawString(230, self.y + 2, note)
            self.y -= 12
        return name

    def checkbox(self, label: str) -> str:
        self._need(26)
        self.n_check += 1
        name = f"CheckBox{self.n_check}"
        self.form.checkbox(name=name, x=50, y=self.y - 13, size=12, buttonStyle="check", borderColor=colors.grey, forceBorder=True, fieldFlags="")
        self.c.setFont("Helvetica", 10)
        self.c.drawString(70, self.y - 11, label)
        self.y -= 26
        return name

    def radio(self, question: str, options: list[str]) -> str:
        self._need(46)
        self.n_group += 1
        name = f"Group{self.n_group}"
        self.c.setFont("Helvetica", 10)
        self.c.drawString(50, self.y - 12, question)
        self.y -= 22
        x = 70
        for opt in options:
            self.form.radio(name=name, value=opt.replace(" ", ""), selected=False, x=x, y=self.y - 13, size=12,
                            buttonStyle="circle", borderColor=colors.grey, forceBorder=True, fieldFlags="noToggleToOff radio")
            self.c.drawString(x + 17, self.y - 11, opt)
            x += 30 + self.c.stringWidth(opt, "Helvetica", 10) + 25
        self.y -= 26
        return name

    def para(self, text: str, size: float = 9) -> None:
        self.c.setFont("Helvetica", size)
        for line in _wrap(text, 95):
            self._need(14)
            self.c.drawString(50, self.y - 10, line)
            self.y -= 13

    def notes_page(self, title: str, notes: list[tuple[str, str]]) -> None:
        self.new_page()
        self.c.setFont("Helvetica-Bold", 14)
        self.c.drawString(50, self.y, title)
        self.y -= 26
        for head, body in notes:
            self.c.setFont("Helvetica-Bold", 10.5)
            self.c.drawString(50, self.y, head)
            self.y -= 15
            self.c.setFont("Helvetica", 10)
            for line in _wrap(body, 92):
                self.c.drawString(50, self.y, line)
                self.y -= 13
            self.y -= 8

    def save(self) -> None:
        self.c.save()


def _wrap(text: str, width: int) -> list[str]:
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        lines.append(cur)
    return lines


def doc_pdf(path: Path, title: str, blocks: list) -> None:
    """Native text PDF via platypus. blocks: str paragraphs, ('h', text), ('table', rows)."""
    story = [Paragraph(title, STYLES["Title"]), Spacer(1, 6)]
    for b in blocks:
        if isinstance(b, str):
            story.append(Paragraph(b, STYLES["BodyText"]))
            story.append(Spacer(1, 4))
        elif b[0] == "h":
            story.append(Paragraph(b[1], STYLES["Heading3"]))
        elif b[0] == "table":
            t = Table(b[1], hAlign="LEFT")
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm).build(story)


def scan(src_pdf: Path, out: Path, seed: int, as_pdf: bool) -> None:
    """Render a native PDF as a degraded grayscale 'scan' (PNG or image-only PDF)."""
    rnd = random.Random(seed)
    doc = fitz.open(src_pdf)
    images = []
    for page in doc:
        pix = page.get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        im = Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")
        im = im.rotate(rnd.uniform(-0.8, 0.8), expand=False, fillcolor=255, resample=Image.BICUBIC)
        px = im.load()
        for _ in range(im.width * im.height // 400):
            x, y = rnd.randrange(im.width), rnd.randrange(im.height)
            px[x, y] = rnd.choice((150, 200, 90))
        im = im.filter(ImageFilter.GaussianBlur(0.6))
        images.append(im)
    if as_pdf:
        out_doc = fitz.open()
        for im in images:
            buf = io.BytesIO()
            im.save(buf, format="PNG")
            p = out_doc.new_page(width=im.width * 72 / 200, height=im.height * 72 / 200)
            p.insert_image(p.rect, stream=buf.getvalue())
        out_doc.save(out)
    else:
        images[0].save(out, dpi=(200, 200))


def reset(name: str) -> Path:
    d = HERE / name
    if d.exists():
        shutil.rmtree(d)
    (d / "sources").mkdir(parents=True)
    return d


def write_gold(d: Path, name: str, description: str, form: str, fields: list[dict]) -> None:
    gold = {"name": name, "description": description, "form": form, "fields": [
        {k: v for k, v in f.items() if k not in ("schema_type", "options", "section", "required")} for f in fields
    ]}
    (d / "gold.json").write_text(json.dumps(gold, indent=2, ensure_ascii=False))
    schema = [
        {
            "label": f["label"],
            "type": f.get("schema_type", "text"),
            **({"options": f["options"]} if f.get("options") else {}),
            **({"section": f["section"]} if f.get("section") else {}),
        }
        for f in fields
    ]
    (d / "expected_schema.json").write_text(json.dumps(schema, indent=2, ensure_ascii=False))


# ============================================================ 01 fillable PDF
def set01() -> None:
    d = reset("01_acroform_basic")
    fb = AcroFormBuilder(d / "form.pdf", "Client Registration Form")
    fb.para("Please complete all sections in BLOCK LETTERS. Fields marked * are required.")
    A, B, C, D = "Section A - Personal details", "Section B - Employment and income", "Section C - Tax residency", "Section D - Declaration"
    fb.heading(A)
    fb.text("1. Full name *")
    fb.text("2. Previous surname (if any)")
    fb.text("3. Date of birth (DD/MM/YYYY) *", width=120)
    fb.text("4. Email address")
    fb.text("5. Mobile phone")
    fb.text("6. Residential address", multiline=True)
    fb.radio("7. Marital status", ["Single", "Married", "De facto", "Other"])
    fb.heading(B)
    fb.text("8. Occupation")
    fb.text("9. Employer ABN", width=160)
    fb.text("10. Annual salary ($)", width=140)
    fb.text("11. Interest income ($)", width=140)
    fb.text("12. Total income ($)", width=140, note="Note: add your salary and interest income.")
    fb.heading(C)
    fb.radio("13. Were you an Australian resident for tax purposes in 2023-24? *", ["Yes", "No"])
    fb.heading(D)
    fb.checkbox("I would like to receive the client newsletter")
    fb.text("Signature of applicant")
    fb.text("Date signed", width=120)
    fb.notes_page("Notes to help you complete this form", [
        ("Notes to question 13",
         "You were an Australian resident for tax purposes in 2023-24 if you were physically present in Australia "
         "for 183 days or more during the income year (1 July 2023 to 30 June 2024). Count every day you were in Australia, "
         "including days of arrival and departure. If you were present for fewer than 183 days, answer No."),
        ("Notes to question 12", "Total income is the sum of your annual salary and your interest income."),
    ])
    fb.save()

    src = d / "sources"
    doc_pdf(src / "employment_letter.pdf", "Harbourline Logistics Pty Ltd", [
        "ABN 51 824 753 556 | Level 4, 100 Collins Street, Melbourne VIC 3000",
        "2 July 2024",
        "To whom it may concern,",
        "This letter confirms that Ms Priya Natarajan has been employed by Harbourline Logistics Pty Ltd on a permanent "
        "full-time basis since 12 February 2019. Her current position is Senior Logistics Coordinator.",
        "Her annual base salary for the 2023-24 financial year was $92,500.00 (before tax), paid fortnightly.",
        "Please contact our HR team on (03) 9012 3456 if you need more information.",
        "Yours sincerely,<br/>Daniel Moreau<br/>HR Manager",
    ])
    doc_pdf(src / "bank_statement.pdf", "Southern Cross Savings Bank - Annual Interest Summary", [
        "Account holder: Priya Natarajan",
        "Mailing address: 27 Wattle Grove, Carlton VIC 3053",
        "Statement period: 1 July 2023 to 30 June 2024",
        ("table", [
            ["Account", "BSB", "Account number", "Interest earned"],
            ["High Interest Saver", "063-182", "1048 2231", "$1,021.15"],
            ["Everyday Access", "063-182", "1048 9910", "$263.22"],
            ["Total interest earned", "", "", "$1,284.37"],
        ]),
        "Withholding tax deducted: $0.00. Keep this summary for your tax return.",
    ])
    doc_pdf(src / "customer_profile.pdf", "Customer Profile", [
        ("h", "Personal details"),
        ("table", [
            ["Field", "Value"],
            ["Name", "Priya Natarajan"],
            ["Date of birth", "14 March 1990"],
            ["Email", "priya.natarajan@example.com"],
            ["Mobile", "0412 345 678"],
            ["Marital status", "Married"],
            ["Residential address", "27 Wattle Grove, Carlton VIC 3053"],
        ]),
        ("h", "Travel record (2023-24 income year)"),
        "According to the passenger movement summary for 1 July 2023 to 30 June 2024, Ms Natarajan was physically present "
        "in Australia for 241 days and overseas for 125 days (secondment to Singapore, 3 Sep 2023 to 5 Jan 2024).",
    ])
    write_gold(d, "01_acroform_basic", "Fillable PDF, native sources, a derived total, a judgment field with a notes page.", "form.pdf", [
        {"label": "Full name", "value": "Priya Natarajan", "schema_type": "text", "section": A},
        {"label": "Previous surname (if any)", "value": None, "schema_type": "text", "section": A},
        {"label": "Date of birth (DD/MM/YYYY)", "value": "1990-03-14", "schema_type": "date", "section": A},
        {"label": "Email address", "value": "priya.natarajan@example.com", "schema_type": "email", "section": A},
        {"label": "Mobile phone", "value": "0412 345 678", "schema_type": "phone", "section": A},
        {"label": "Residential address", "value": "27 Wattle Grove, Carlton VIC 3053", "schema_type": "multiline", "section": A},
        {"label": "Marital status", "value": "Married", "schema_type": "choice", "options": ["Single", "Married", "De facto", "Other"], "section": A},
        {"label": "Occupation", "value": "Senior Logistics Coordinator", "schema_type": "text", "section": B},
        {"label": "Employer ABN", "value": "51 824 753 556", "schema_type": "text", "section": B},
        {"label": "Annual salary ($)", "value": "92500.00", "schema_type": "currency", "section": B},
        {"label": "Interest income ($)", "value": "1284.37", "schema_type": "currency", "section": B},
        {"label": "Total income ($)", "value": "93784.37", "derived": True, "schema_type": "currency", "section": B},
        {"label": "Were you an Australian resident for tax purposes in 2023-24?", "kind": "judgment", "value": "Yes",
         "facts": {"days present in Australia during 2023-24": "241"}, "schema_type": "choice", "options": ["Yes", "No"], "section": C},
        {"label": "I would like to receive the client newsletter", "value": None, "schema_type": "checkbox", "section": D},
        {"label": "Signature of applicant", "value": None, "skip_eval": True, "schema_type": "signature", "section": D},
        {"label": "Date signed", "value": None, "schema_type": "date", "section": D},
    ])


# ============================================================ 02 Word form
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _sdt_xml(alias: str, kind: str, placeholder: str, items: list[str] | None = None, sid: int = 1) -> str:
    pr = f'<w:alias w:val="{alias}"/><w:tag w:val="{alias.lower().replace(" ", "_")}"/><w:id w:val="{sid}"/><w:showingPlcHdr/>'
    if kind == "date":
        pr += '<w:date><w:dateFormat w:val="dd/MM/yyyy"/><w:lid w:val="en-AU"/></w:date>'
    elif kind == "dropdown":
        pr += "<w:dropDownList>" + "".join(f'<w:listItem w:displayText="{i}" w:value="{i}"/>' for i in items or []) + "</w:dropDownList>"
    elif kind == "checkbox":
        pr = pr.replace("<w:showingPlcHdr/>", "")
        pr += ('<w14:checkbox><w14:checked w14:val="0"/><w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
               '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox>')
        placeholder = "☐"
    else:
        pr += "<w:text/>"
    rstyle = "" if kind == "checkbox" else '<w:rPr><w:rStyle w:val="PlaceholderText"/></w:rPr>'
    return (f'<w:sdt xmlns:w="{WNS}" xmlns:w14="{W14}"><w:sdtPr>{pr}</w:sdtPr>'
            f'<w:sdtContent><w:r>{rstyle}<w:t xml:space="preserve">{placeholder}</w:t></w:r></w:sdtContent></w:sdt>')


def set02() -> None:
    d = reset("02_docx_form")
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Volunteer Application Form", 0)
    doc.add_paragraph("Thank you for your interest in volunteering with Riverbend Community Services.")
    doc.add_heading("Personal details", 1)
    for i, (alias, kind, items) in enumerate([
        ("Full name", "text", None),
        ("Date of birth", "date", None),
        ("Preferred contact method", "dropdown", ["Email", "Phone", "Post"]),
    ]):
        p = doc.add_paragraph()
        r = p.add_run(f"{alias}: ")
        r.bold = True
        p._p.append(parse_xml(_sdt_xml(alias, kind, "Click or tap here to enter text.", items, sid=100 + i)))
    doc.add_heading("Contact details", 1)
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    for row, label in zip(t.rows, ["Email", "Phone", "Suburb", "Postcode"]):
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
    doc.add_heading("Availability", 1)
    doc.add_paragraph("Days available per week: ________")
    doc.add_paragraph("Emergency contact name: ________________   Relationship: ____________")
    doc.add_paragraph("Do you hold a current driver licence?   ☐ Yes   ☐ No")
    p = doc.add_paragraph()
    p._p.append(parse_xml(_sdt_xml("Police check consent", "checkbox", "", sid=200)))
    p.add_run(" I consent to a national police check")
    doc.add_heading("Declaration", 1)
    doc.add_paragraph("Signature: ______________________   Date: ____________")
    doc.save(d / "form.docx")

    src = d / "sources"
    cv = docx.Document()
    cv.add_heading("Amara Okafor", 0)
    cv.add_paragraph("Email: amara.okafor@example.org | Phone: 0433 908 172")
    cv.add_paragraph("Address: 5/18 Riverside Parade, Abbotsford VIC 3067")
    cv.add_paragraph("Date of birth: 07/11/1996")
    cv.add_heading("Availability and preferences", 1)
    cv.add_paragraph("I am available 3 days per week (Monday, Wednesday and Saturday). Please contact me by email.")
    cv.add_paragraph("Licences: Current full Victorian driver licence (expires 2029).")
    cv.add_heading("Experience", 1)
    tbl = cv.add_table(rows=3, cols=3)
    tbl.style = "Table Grid"
    for row, vals in zip(tbl.rows, [["Organisation", "Role", "Years"], ["FoodShare Melbourne", "Warehouse volunteer", "2021-2023"], ["Abbotsford Library", "Homework club tutor", "2023-present"]]):
        for c, v in zip(row.cells, vals):
            c.text = v
    cv.add_heading("Emergency contact", 1)
    cv.add_paragraph("Emergency contact: Chidi Okafor (brother), 0402 557 310")
    cv.save(src / "resume.docx")
    doc_pdf(src / "reference_letter.pdf", "Abbotsford Library - Volunteer Reference", [
        "To the volunteer coordinator,",
        "Amara Okafor has volunteered as a homework club tutor at Abbotsford Library since March 2023. "
        "She is reliable, punctual and works well with children and families.",
        "Kind regards,<br/>Helen Brooks, Branch Manager",
    ])
    write_gold(d, "02_docx_form", "Word form: content controls, label/value table, blank lines, checkbox glyphs; DOCX source.", "form.docx", [
        {"label": "Full name", "value": "Amara Okafor", "schema_type": "text", "section": "Personal details"},
        {"label": "Date of birth", "value": "1996-11-07", "schema_type": "date", "section": "Personal details"},
        {"label": "Preferred contact method", "value": "Email", "schema_type": "choice", "options": ["Email", "Phone", "Post"], "section": "Personal details"},
        {"label": "Email", "value": "amara.okafor@example.org", "schema_type": "email", "section": "Contact details"},
        {"label": "Phone", "value": "0433 908 172", "schema_type": "phone", "section": "Contact details"},
        {"label": "Suburb", "value": "Abbotsford", "schema_type": "text", "section": "Contact details"},
        {"label": "Postcode", "value": "3067", "schema_type": "text", "section": "Contact details"},
        {"label": "Days available per week", "value": "3", "schema_type": "number", "section": "Availability"},
        {"label": "Emergency contact name", "value": "Chidi Okafor", "schema_type": "text", "section": "Availability"},
        {"label": "Relationship", "value": "Brother", "schema_type": "text", "section": "Availability"},
        {"label": "Do you hold a current driver licence?", "value": "Yes", "schema_type": "choice", "options": ["Yes", "No"], "section": "Availability"},
        {"label": "Police check consent", "value": None, "schema_type": "checkbox", "section": "Availability"},
        {"label": "Signature", "value": None, "skip_eval": True, "schema_type": "signature", "section": "Declaration"},
        {"label": "Date", "value": None, "schema_type": "date", "section": "Declaration"},
    ])


# ============================================================ 03 scanned sources
def set03() -> None:
    d = reset("03_scanned_source")
    fb = AcroFormBuilder(d / "form.pdf", "Notice of Disposal - Vehicle Transfer")
    S1, S2 = "Buyer details", "Vehicle and sale details"
    fb.heading(S1)
    fb.text("Buyer full name")
    fb.text("Driver licence number", width=160)
    fb.text("Licence expiry date", width=120)
    fb.text("Buyer address", multiline=True)
    fb.heading(S2)
    fb.text("Registration number", width=120)
    fb.text("VIN", width=220)
    fb.text("Sale price ($)", width=120)
    fb.text("Date of sale", width=120)
    fb.text("Odometer reading (km)", width=120)
    fb.save()
    tmp = d / "_tmp"
    tmp.mkdir()
    doc_pdf(tmp / "licence.pdf", "DRIVER LICENCE", [
        "STATE OF NEW SOUTH WALES (SPECIMEN)",
        ("table", [["Name", "MARCUS JAMES TAYLOR"], ["Address", "41 Kurrajong Road, Wagga Wagga NSW 2650"],
                   ["Licence No", "2847 1956"], ["Date of birth", "09 MAY 1984"], ["Expiry", "22 NOV 2029"], ["Class", "C"]]),
    ])
    doc_pdf(tmp / "receipt.pdf", "Private Vehicle Sale - Receipt", [
        "Seller: Angela Ruiz, 7 Tallowwood Close, Albury NSW 2640",
        "Buyer: Marcus James Taylor",
        ("table", [["Item", "Detail"], ["Vehicle", "2013 Toyota Corolla Ascent"], ["Registration", "XKD-482"],
                   ["VIN", "6T1BF3EK7DX123456"], ["Odometer", "148,920 km"], ["Agreed price", "$18,750.00"]]),
        "Date of sale: 3 August 2024. Paid in full by bank transfer.",
        "Signed by both parties.",
    ])
    scan(tmp / "licence.pdf", d / "sources" / "licence_scan.png", seed=3, as_pdf=False)
    scan(tmp / "receipt.pdf", d / "sources" / "receipt_scan.pdf", seed=4, as_pdf=True)
    shutil.rmtree(tmp)
    write_gold(d, "03_scanned_source", "Fillable PDF with scanned (image-only) sources that need OCR.", "form.pdf", [
        {"label": "Buyer full name", "value": "Marcus James Taylor", "schema_type": "text", "section": S1},
        {"label": "Driver licence number", "value": "2847 1956", "schema_type": "text", "section": S1},
        {"label": "Licence expiry date", "value": "2029-11-22", "schema_type": "date", "section": S1},
        {"label": "Buyer address", "value": "41 Kurrajong Road, Wagga Wagga NSW 2650", "schema_type": "multiline", "section": S1},
        {"label": "Registration number", "value": "XKD-482", "schema_type": "text", "section": S2},
        {"label": "VIN", "value": "6T1BF3EK7DX123456", "schema_type": "text", "section": S2},
        {"label": "Sale price ($)", "value": "18750.00", "schema_type": "currency", "section": S2},
        {"label": "Date of sale", "value": "2024-08-03", "schema_type": "date", "section": S2},
        {"label": "Odometer reading (km)", "value": "148920", "schema_type": "number", "section": S2},
    ])


# ============================================================ 04 conflicting sources
def set04() -> None:
    d = reset("04_conflicting_sources")
    fb = AcroFormBuilder(d / "form.pdf", "Change of Details Notification")
    S = "Your details"
    fb.heading(S)
    fb.text("Full name")
    fb.text("Customer number", width=140)
    fb.text("Date of birth", width=120)
    fb.text("Current residential address", multiline=True)
    fb.text("Phone number", width=160)
    fb.text("Email address")
    fb.save()
    src = d / "sources"
    doc_pdf(src / "energy_bill_march_2024.pdf", "Brightwave Energy - Tax Invoice", [
        "Customer: Tomasz Kowalczyk | Customer number: BW-5531-208",
        "Supply address: 8 Banksia Street, Fitzroy VIC 3065",
        "Phone on file: 0400 111 222",
        "Invoice date: 14 March 2024 | Amount due: $212.40",
    ])
    doc_pdf(src / "council_rates_notice_2024.pdf", "City of Merri-bek - Rates Notice 2024-25", [
        "Owner/occupier: Tomasz Kowalczyk",
        "Property address: 15/220 Lygon Street, Brunswick East VIC 3057",
        "Contact phone: 0400 111 222. Date of birth (concession check): 30/06/1979",
        "Notice issued: 22 August 2024",
    ])
    write_gold(d, "04_conflicting_sources", "Two sources give different addresses; one phone agrees; email absent.", "form.pdf", [
        {"label": "Full name", "value": "Tomasz Kowalczyk", "schema_type": "text", "section": S},
        {"label": "Customer number", "value": "BW-5531-208", "schema_type": "text", "section": S},
        {"label": "Date of birth", "value": "1979-06-30", "schema_type": "date", "section": S},
        {"label": "Current residential address", "conflict": ["8 Banksia Street, Fitzroy VIC 3065", "15/220 Lygon Street, Brunswick East VIC 3057"], "schema_type": "multiline", "section": S},
        {"label": "Phone number", "value": "0400 111 222", "schema_type": "phone", "section": S},
        {"label": "Email address", "value": None, "schema_type": "email", "section": S},
    ])


# ============================================================ 05 absent fields
def set05() -> None:
    d = reset("05_absent_fields")
    fb = AcroFormBuilder(d / "form.pdf", "Motor Insurance Claim Form")
    S1, S2, S3, S4 = "Policy details", "Incident details", "Witness and police", "Payment details"
    fb.heading(S1)
    fb.text("Policy number", width=160)
    fb.text("Policyholder name")
    fb.text("Other insurer (if any)")
    fb.heading(S2)
    fb.text("Date of incident", width=120)
    fb.text("Location of incident")
    fb.text("Estimated claim amount ($)", width=140)
    fb.text("Injury description (if any)", multiline=True)
    fb.text("Hospital attended")
    fb.heading(S3)
    fb.text("Witness name")
    fb.text("Witness phone", width=160)
    fb.text("Police report number", width=160)
    fb.heading(S4)
    fb.text("BSB", width=90)
    fb.text("Account number", width=140)
    fb.text("Number of previous claims", width=60)
    fb.save()
    doc_pdf(d / "sources" / "incident_email.pdf", "Email: Claim lodgement", [
        "From: Sofia Lindqvist &lt;sofia.l@example.net&gt;<br/>To: claims@harbourmutual.example<br/>Date: 19 May 2024",
        "Hi, I'd like to lodge a claim on my policy HM-77310492. On 17 May 2024 another car reversed into my parked "
        "vehicle in the Coles car park on Sydney Road, Coburg. Nobody was hurt and no police report was made. "
        "The panel beater quoted $2,340.00 for repairs.",
        "I didn't get the other driver's details. You can call me back any time.",
        "Thanks, Sofia Lindqvist",
        "-- Harbour Mutual Claims: call 1300 555 019 (8am-8pm) --",
    ])
    write_gold(d, "05_absent_fields", "Most fields are absent from the sources; distractor numbers present.", "form.pdf", [
        {"label": "Policy number", "value": "HM-77310492", "schema_type": "text", "section": S1},
        {"label": "Policyholder name", "value": "Sofia Lindqvist", "schema_type": "text", "section": S1},
        {"label": "Other insurer (if any)", "value": None, "schema_type": "text", "section": S1},
        {"label": "Date of incident", "value": "2024-05-17", "schema_type": "date", "section": S2},
        {"label": "Location of incident", "value": "Coles car park on Sydney Road, Coburg", "schema_type": "text", "section": S2},
        {"label": "Estimated claim amount ($)", "value": "2340.00", "schema_type": "currency", "section": S2},
        {"label": "Injury description (if any)", "value": None, "schema_type": "multiline", "section": S2},
        {"label": "Hospital attended", "value": None, "schema_type": "text", "section": S2},
        {"label": "Witness name", "value": None, "schema_type": "text", "section": S3},
        {"label": "Witness phone", "value": None, "schema_type": "phone", "section": S3},
        {"label": "Police report number", "value": None, "schema_type": "text", "section": S3},
        {"label": "BSB", "value": None, "schema_type": "text", "section": S4},
        {"label": "Account number", "value": None, "schema_type": "text", "section": S4},
        {"label": "Number of previous claims", "value": None, "schema_type": "number", "section": S4},
    ])


# ============================================================ 06 judgment
def set06() -> None:
    d = reset("06_judgment")
    fb = AcroFormBuilder(d / "form.pdf", "Household Concessions Application")
    S1, S2 = "Part A - Applicant", "Part B - Eligibility"
    fb.heading(S1)
    fb.text("1. Full name")
    fb.text("2. Date of birth", width=120)
    fb.text("3. Pensioner Concession Card number", width=180)
    fb.heading(S2)
    fb.radio("4. Were you aged 65 or over on 1 January 2025?", ["Yes", "No"])
    fb.radio("5. Do you hold a current Pensioner Concession Card?", ["Yes", "No"])
    fb.radio("6. Are you eligible for the energy rebate?", ["Yes", "No"])
    fb.radio("7. Do you qualify for the water concession?", ["Yes", "No"])
    fb.notes_page("Notes", [
        ("Notes to question 4", "Answer Yes if your 65th birthday was on or before 1 January 2025."),
        ("Notes to question 5", "Your card is current if its expiry date is after the date you sign this form (any date in 2025)."),
        ("Notes to question 6",
         "You are eligible for the energy rebate only if ALL of these apply: (a) you hold a current Pensioner Concession Card; "
         "(b) the electricity account is in your own name; and (c) the account is for the home you live in. "
         "If the account is in someone else's name, answer No."),
        ("Notes to question 7",
         "You qualify for the water concession if you own your home or you are a tenant who pays water usage charges, "
         "AND you have not received the water concession at any address in the last 12 months."),
    ])
    fb.save()
    src = d / "sources"
    doc_pdf(src / "concession_card_letter.pdf", "Services Australia (specimen) - Your Pensioner Concession Card", [
        "Mrs Mei Chen<br/>3 Jacaranda Court, Box Hill VIC 3128",
        "Date: 5 October 2024",
        "Your Pensioner Concession Card number is 312 456 789A. Your card is valid until 30/09/2026.",
        "Our records show your date of birth as 2 February 1958.",
    ])
    doc_pdf(src / "electricity_bill.pdf", "Volta Power - Electricity Bill", [
        "Account holder: Robert Chen",
        "Supply address: 3 Jacaranda Court, Box Hill VIC 3128",
        "Account number: 7781 0044 19 | Billing period: 1 Jul 2024 - 30 Sep 2024",
        "Total due: $388.15 by 28 October 2024",
    ])
    write_gold(d, "06_judgment", "Judgment fields: clear yes, reasoning-to-no, and a cannot_determine with missing facts.", "form.pdf", [
        {"label": "Full name", "value": "Mei Chen", "schema_type": "text", "section": S1},
        {"label": "Date of birth", "value": "1958-02-02", "schema_type": "date", "section": S1},
        {"label": "Pensioner Concession Card number", "value": "312 456 789A", "schema_type": "text", "section": S1},
        {"label": "Were you aged 65 or over on 1 January 2025?", "kind": "judgment", "value": "Yes",
         "facts": {"date of birth": "1958-02-02"}, "schema_type": "choice", "options": ["Yes", "No"], "section": S2},
        {"label": "Do you hold a current Pensioner Concession Card?", "kind": "judgment", "value": "Yes",
         "facts": {"card expiry date": "2026-09-30"}, "schema_type": "choice", "options": ["Yes", "No"], "section": S2},
        {"label": "Are you eligible for the energy rebate?", "kind": "judgment", "value": "No",
         "facts": {"electricity account holder": "Robert Chen", "holds concession card": "yes"},
         "schema_type": "choice", "options": ["Yes", "No"], "section": S2},
        {"label": "Do you qualify for the water concession?", "kind": "judgment", "value": "cannot_determine",
         "missing_facts": ["owns home or tenant paying water charges", "received water concession in the last 12 months"],
         "user_answers": [
             {"match": ["own", "tenant", "home", "water charges", "ownership"], "value": "Yes"},
             {"match": ["received", "last 12 months", "previous", "already"], "value": "No"},
         ],
         "value_after_answers": "Yes",
         "schema_type": "choice", "options": ["Yes", "No"], "section": S2},
    ])


# ============================================================ 07 flat PDF
def set07() -> None:
    d = reset("07_flat_pdf")
    c = canvas.Canvas(str(d / "form.pdf"), pagesize=A4)
    y = H - 60
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, "Library Membership Application")
    y -= 34
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "Member details")
    y -= 30
    c.setFont("Helvetica", 10)
    for label in ["Full name", "Date of birth", "Street address", "Suburb and postcode", "Phone", "Email"]:
        c.drawString(50, y, label + ":")
        c.setLineWidth(0.6)
        c.line(170, y - 2, 520, y - 2)
        y -= 30
    y -= 6
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "Preferences")
    y -= 28
    c.setFont("Helvetica", 10)
    c.rect(50, y - 2, 11, 11)
    c.drawString(68, y, "Send me the monthly events newsletter")
    y -= 26
    c.drawString(50, y, "Preferred branch: ____________________________")
    y -= 40
    c.drawString(50, y, "Signature: ______________________          Date: ______________")
    c.save()
    doc_pdf(d / "sources" / "proof_of_address.pdf", "Glenfern Water - Account Summary", [
        "Customer: Lucas Ferreira<br/>12 Moonah Avenue, Glenroy VIC 3046",
        "Customer phone: 0491 220 846 | Email: lucas.ferreira@example.com",
        "Date of birth (identity check): 21 January 2001",
        "Account summary issued 2 September 2024.",
    ])
    doc_pdf(d / "sources" / "note.pdf", "Note from applicant", [
        "I'd like to join the Glenroy branch. Please don't send me newsletters, thanks.",
    ])
    S1, S2 = "Member details", "Preferences"
    write_gold(d, "07_flat_pdf", "Flat PDF with no form fields: underscores, drawn lines and a box.", "form.pdf", [
        {"label": "Full name", "value": "Lucas Ferreira", "schema_type": "text", "section": S1},
        {"label": "Date of birth", "value": "2001-01-21", "schema_type": "date", "section": S1},
        {"label": "Street address", "value": "12 Moonah Avenue", "schema_type": "text", "section": S1},
        {"label": "Suburb and postcode", "value": "Glenroy VIC 3046", "schema_type": "text", "section": S1},
        {"label": "Phone", "value": "0491 220 846", "schema_type": "phone", "section": S1},
        {"label": "Email", "value": "lucas.ferreira@example.com", "schema_type": "email", "section": S1},
        {"label": "Send me the monthly events newsletter", "value": False, "schema_type": "checkbox", "section": S2},
        {"label": "Preferred branch", "value": "Glenroy", "schema_type": "text", "section": S2},
        {"label": "Signature", "value": None, "skip_eval": True, "schema_type": "signature", "section": S2},
        {"label": "Date", "value": None, "schema_type": "date", "section": S2},
    ])


if __name__ == "__main__":
    for fn in (set01, set02, set03, set04, set05, set06, set07):
        fn()
        print("built", fn.__name__)
